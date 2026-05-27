# agent/tools.py — Herramientas del agente Torre Fuerte
# Generado por AgentKit

import os
import re
import yaml
import logging
import httpx
from datetime import datetime, timedelta
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

logger = logging.getLogger("agentkit")

_COL = timedelta(hours=-5)


def _ahora_colombia() -> str:
    """Retorna la fecha y hora actual en Colombia (UTC-5)."""
    return (datetime.utcnow() + _COL).strftime("%Y-%m-%d %H:%M")


RESEND_API_KEY = os.getenv("RESEND_API_KEY", "")
EMAIL_LEADS = os.getenv("EMAIL_LEADS", "")

KNOWLEDGE_DIR = "knowledge"
EXTENSIONES_MEDIA = {".jpg", ".jpeg", ".png", ".gif", ".mp4", ".pdf"}

# ── Helpers de configuración ──────────────────────────────────────────────────

def cargar_info_negocio() -> dict:
    """Carga la configuración completa desde config/business.yaml."""
    try:
        with open("config/business.yaml", "r", encoding="utf-8") as f:
            return yaml.safe_load(f) or {}
    except FileNotFoundError:
        logger.error("config/business.yaml no encontrado")
        return {}


def _cfg() -> dict:
    """Alias corto de cargar_info_negocio para uso interno."""
    return cargar_info_negocio()


def _nombre_corto() -> str:
    """Retorna el nombre corto del negocio para usar en emails y mensajes."""
    cfg = _cfg()
    return cfg.get("negocio", {}).get("nombre_corto") or cfg.get("negocio", {}).get("nombre", "")


def _normalizar_codigo(codigo: str) -> str:
    """Normaliza un código de apartamento: minúsculas, sin espacios ni guiones extra."""
    return codigo.lower().replace("apto", "").replace("-", "").replace(" ", "").strip()


def obtener_plano(codigo_apto: str) -> str | None:
    """Retorna la ruta al PDF de plano dado un código de apartamento.
    Lee el mapeo desde config/business.yaml → planos."""
    codigo = _normalizar_codigo(codigo_apto)
    mapa = _cfg().get("planos", {})
    archivo = mapa.get(codigo)
    if not archivo:
        return None
    ruta = os.path.join(KNOWLEDGE_DIR, "planos", archivo)
    return ruta if os.path.exists(ruta) else None


def _grupo_renders(clave_norm: str) -> dict | None:
    """Encuentra el grupo de renders que corresponde a la clave normalizada."""
    for grupo in _cfg().get("renders", []):
        claves_norm = [_normalizar_codigo(c) for c in grupo.get("claves", [])]
        if clave_norm in claves_norm:
            return grupo
    return None


def obtener_renders(clave: str) -> list[str]:
    """Retorna rutas locales de renders (para test_local.py o si no hay BASE_URL)."""
    grupo = _grupo_renders(_normalizar_codigo(clave))
    if not grupo:
        return []
    carpeta = os.path.join(KNOWLEDGE_DIR, grupo["carpeta"])
    if not os.path.isdir(carpeta):
        return []
    archivos = sorted([
        os.path.join(carpeta, f)
        for f in os.listdir(carpeta)
        if not f.startswith(".") and os.path.splitext(f)[1].lower() in EXTENSIONES_MEDIA
    ])
    return archivos


def obtener_urls_renders(clave: str, base_url: str = "") -> list[str]:
    """Retorna URLs públicas de renders servidas desde Railway.
    Las rutas se leen de config/business.yaml → renders."""
    grupo = _grupo_renders(_normalizar_codigo(clave))
    if not grupo:
        return []
    carpeta = grupo["carpeta"]
    urls = []
    for img in grupo.get("imagenes", []):
        urls.append(f"{base_url}/knowledge/{carpeta}/{img}")
    video = grupo.get("video")
    if video:
        urls.append(f"{base_url}/knowledge/{carpeta}/{video}")
    return urls


def extraer_marcadores_plano(texto: str) -> tuple[str, list[str]]:
    """Extrae marcadores [PLANO:XXX] del texto de Claude."""
    patron = re.compile(r'\[PLANO:([^\]]+)\]', re.IGNORECASE)
    codigos = [m.group(1).strip() for m in patron.finditer(texto)]
    texto_limpio = patron.sub("", texto).strip()
    return texto_limpio, codigos


def extraer_marcador_lead(texto: str) -> tuple[str, dict | None]:
    """Extrae marcador [LEAD:nombre|email|apto|habitaciones|temperatura|intencion]."""
    patron = re.compile(r'\[LEAD:([^\]]*)\]', re.IGNORECASE)
    m = patron.search(texto)
    if not m:
        return texto, None
    partes = [p.strip() for p in m.group(1).split("|")]
    lead = {
        "nombre": partes[0] if len(partes) > 0 else "",
        "email": partes[1] if len(partes) > 1 else "",
        "apto": partes[2] if len(partes) > 2 else "",
        "habitaciones": partes[3] if len(partes) > 3 else "",
        "temperatura": partes[4] if len(partes) > 4 else "",
        "intencion": partes[5] if len(partes) > 5 else "",
    }
    texto_limpio = patron.sub("", texto).strip()
    return texto_limpio, lead if lead["nombre"] else None


ICONOS_TEMPERATURA = {"caliente": "🔥", "tibio": "🌡️", "frío": "❄️"}


def enviar_email_lead(telefono: str, nombre: str, email: str = "", apto: str = "", habitaciones: str = "", temperatura: str = "", intencion: str = "") -> bool:
    """Envía notificación de nuevo lead via Resend API."""
    if not all([RESEND_API_KEY, EMAIL_LEADS]):
        logger.warning("RESEND_API_KEY o EMAIL_LEADS no configurados")
        return False
    try:
        negocio = _nombre_corto()
        icono = ICONOS_TEMPERATURA.get(temperatura.lower(), "")
        temp_texto = f"{icono} {temperatura.upper()}" if temperatura else "No determinada"
        cuerpo = (
            f"Nuevo lead interesado en {negocio}\n\n"
            f"Nombre:         {nombre}\n"
            f"Teléfono:       {telefono}\n"
            f"Email:          {email or 'No proporcionado'}\n"
            f"Apto interés:   {apto or 'No especificado'}\n"
            f"Habitaciones:   {habitaciones or 'No especificado'}\n"
            f"Intención:      {intencion or 'No especificada'}\n"
            f"Temperatura:    {temp_texto}\n"
            f"Fecha:          {_ahora_colombia()}\n"
        )
        r = httpx.post(
            "https://api.resend.com/emails",
            headers={"Authorization": f"Bearer {RESEND_API_KEY}"},
            json={
                "from": f"{negocio} <onboarding@resend.dev>",
                "to": [EMAIL_LEADS],
                "subject": f"[{temp_texto}] Nuevo lead {negocio} — {nombre}",
                "text": cuerpo,
            },
            timeout=15,
        )
        if r.status_code == 200:
            logger.info(f"Email de lead enviado: {nombre} ({telefono})")
            return True
        logger.error(f"Error Resend: {r.status_code} — {r.text}")
        return False
    except Exception as e:
        logger.error(f"Error enviando email de lead: {e}")
        return False


def extraer_marcadores_render(texto: str) -> tuple[str, list[str]]:
    """Extrae marcadores [RENDER:XXX] del texto de Claude."""
    patron = re.compile(r'\[RENDER:([^\]]+)\]', re.IGNORECASE)
    claves = [m.group(1).strip() for m in patron.finditer(texto)]
    texto_limpio = patron.sub("", texto).strip()
    return texto_limpio, claves


def extraer_marcador_handoff(texto: str) -> tuple[str, str | None]:
    """Extrae [HANDOFF] o [HANDOFF:razon] del texto. Retorna (texto_limpio, razon o None)."""
    patron = re.compile(r'\[HANDOFF(?::([^\]]*))?\]', re.IGNORECASE)
    m = patron.search(texto)
    if not m:
        return texto, None
    razon = (m.group(1) or "conversación completada").strip()
    texto_limpio = patron.sub("", texto).strip()
    return texto_limpio, razon


def extraer_marcador_reagendar_visita(texto: str) -> tuple[str, dict | None]:
    """Extrae [REAGENDAR_VISITA:id|YYYY-MM-DD|HH:MM] del texto.

    Retorna (texto_limpio, dict con id/fecha/hora) o (texto, None).
    """
    patron = re.compile(r'\[REAGENDAR_VISITA:([^\]]+)\]', re.IGNORECASE)
    m = patron.search(texto)
    if not m:
        return texto, None
    partes = [p.strip() for p in m.group(1).split("|")]
    if len(partes) < 3:
        return patron.sub("", texto).strip(), None
    try:
        visita_id = int(partes[0])
    except ValueError:
        return patron.sub("", texto).strip(), None
    texto_limpio = patron.sub("", texto).strip()
    return texto_limpio, {"id": visita_id, "fecha": partes[1], "hora": partes[2]}


def extraer_marcador_cancelar_visita(texto: str) -> tuple[str, int | None]:
    """Extrae [CANCELAR_VISITA:id] del texto. Retorna (texto_limpio, visita_id o None)."""
    patron = re.compile(r'\[CANCELAR_VISITA:(\d+)\]', re.IGNORECASE)
    m = patron.search(texto)
    if not m:
        return texto, None
    visita_id = int(m.group(1))
    texto_limpio = patron.sub("", texto).strip()
    return texto_limpio, visita_id


def extraer_marcador_visita(texto: str) -> tuple[str, dict | None]:
    """Extrae [VISITA:nombre|YYYY-MM-DD|HH:MM|notas] del texto.

    Retorna (texto_limpio, dict con keys nombre/fecha/hora/notas) o (texto, None).
    """
    patron = re.compile(r'\[VISITA:([^\]]+)\]', re.IGNORECASE)
    m = patron.search(texto)
    if not m:
        return texto, None
    partes = [p.strip() for p in m.group(1).split("|")]
    if len(partes) < 3:
        return patron.sub("", texto).strip(), None
    nombre = partes[0] or "Lead"
    fecha = partes[1]
    hora = partes[2]
    notas = partes[3] if len(partes) > 3 else ""
    texto_limpio = patron.sub("", texto).strip()
    return texto_limpio, {"nombre": nombre, "fecha": fecha, "hora": hora, "notas": notas}


def enviar_email_handoff(
    telefono: str,
    nombre: str,
    temperatura: str,
    razon: str,
    apto: str = "",
    habitaciones: str = "",
    email: str = "",
    intencion: str = "",
    resumen: str = "",
) -> bool:
    """Notifica al asesor por email que debe tomar esta conversación en Meta Business Suite."""
    if not all([RESEND_API_KEY, EMAIL_LEADS]):
        logger.warning("RESEND_API_KEY o EMAIL_LEADS no configurados — no se envió email handoff")
        return False
    try:
        icono = ICONOS_TEMPERATURA.get(temperatura.lower(), "🔔")
        temp_texto = f"{icono} {temperatura.upper()}" if temperatura else "No determinada"
        tel_limpio = telefono.lstrip("+")
        negocio = _nombre_corto()
        cuerpo_html = f"""
<div style="font-family:Arial,sans-serif;max-width:600px;margin:0 auto">
  <div style="background:#1a3c5e;padding:20px;border-radius:8px 8px 0 0">
    <h2 style="color:#fff;margin:0">🔔 Transferencia a Asesor</h2>
    <p style="color:#aac8e4;margin:4px 0 0">{negocio}</p>
  </div>
  <div style="background:#f5f8fb;padding:24px;border-radius:0 0 8px 8px;border:1px solid #dce8f3">
    <table style="width:100%;border-collapse:collapse">
      <tr><td style="padding:8px 0;color:#666;width:140px">📱 Teléfono</td>
          <td style="padding:8px 0;font-weight:bold;font-size:18px">+{tel_limpio}</td></tr>
      <tr><td style="padding:8px 0;color:#666">👤 Nombre</td>
          <td style="padding:8px 0">{nombre}</td></tr>
      <tr><td style="padding:8px 0;color:#666">📧 Email</td>
          <td style="padding:8px 0">{email or 'No proporcionado'}</td></tr>
      <tr><td style="padding:8px 0;color:#666">🏠 Apto interés</td>
          <td style="padding:8px 0">{apto or 'No especificado'}</td></tr>
      <tr><td style="padding:8px 0;color:#666">🛏️ Habitaciones</td>
          <td style="padding:8px 0">{habitaciones or 'No especificado'}</td></tr>
      <tr><td style="padding:8px 0;color:#666">💼 Intención</td>
          <td style="padding:8px 0">{intencion.capitalize() if intencion else 'No especificada'}</td></tr>
      <tr><td style="padding:8px 0;color:#666">🌡️ Temperatura</td>
          <td style="padding:8px 0"><strong>{temp_texto}</strong></td></tr>
      <tr><td style="padding:8px 0;color:#666">📋 Razón</td>
          <td style="padding:8px 0">{razon}</td></tr>
      <tr><td style="padding:8px 0;color:#666">🕐 Hora</td>
          <td style="padding:8px 0">{_ahora_colombia()}</td></tr>
    </table>
    {f'''<div style="margin-top:20px;background:#fff8e1;border-left:4px solid #f39c12;padding:14px 16px;border-radius:4px">
      <div style="font-size:12px;font-weight:bold;color:#d68910;margin-bottom:8px">📋 RESUMEN DE LA CONVERSACIÓN</div>
      <div style="font-size:14px;color:#555;white-space:pre-line">{resumen}</div>
    </div>''' if resumen else ''}
    <div style="margin-top:24px;text-align:center">
      <a href="https://business.facebook.com/latest/inbox/all"
         style="display:inline-block;background:#1a3c5e;color:#fff;padding:12px 28px;
                border-radius:6px;text-decoration:none;font-weight:bold;margin:0 8px">
        Abrir Meta Business Suite
      </a>
    </div>
    <p style="margin-top:16px;color:#888;font-size:13px;text-align:center">
      En Meta Business Suite, busca el número <strong>+{tel_limpio}</strong> en la bandeja de entrada.
    </p>
  </div>
</div>
"""
        cuerpo_texto = (
            f"TRANSFERENCIA A ASESOR — {negocio}\n\n"
            f"Teléfono:     +{tel_limpio}\n"
            f"Nombre:       {nombre}\n"
            f"Email:        {email or 'No proporcionado'}\n"
            f"Apto interés: {apto or 'No especificado'}\n"
            f"Habitaciones: {habitaciones or 'No especificado'}\n"
            f"Intención:    {intencion.capitalize() if intencion else 'No especificada'}\n"
            f"Temperatura:  {temp_texto}\n"
            f"Razón:        {razon}\n"
            f"Hora:         {_ahora_colombia()}\n"
            + (f"\nResumen:\n{resumen}\n" if resumen else "")
            + f"\nAbre Meta Business Suite y busca +{tel_limpio}:\n"
            f"https://business.facebook.com/latest/inbox/all\n"
        )
        r = httpx.post(
            "https://api.resend.com/emails",
            headers={"Authorization": f"Bearer {RESEND_API_KEY}"},
            json={
                "from": f"{negocio} <onboarding@resend.dev>",
                "to": [EMAIL_LEADS],
                "subject": f"🔔 [{temp_texto}] {nombre} — +{tel_limpio}",
                "html": cuerpo_html,
                "text": cuerpo_texto,
            },
            timeout=15,
        )
        if r.status_code == 200:
            logger.info(f"Email handoff enviado: {nombre} ({telefono})")
            return True
        logger.error(f"Error Resend handoff: {r.status_code} — {r.text}")
        return False
    except Exception as e:
        logger.error(f"Error enviando email handoff: {e}")
        return False


def obtener_disponibilidad() -> list[dict]:
    """Retorna la lista de apartamentos desde config/business.yaml → disponibilidad."""
    return _cfg().get("disponibilidad", [])


def registrar_lead(telefono: str, nombre: str, interes: str, presupuesto: str = "") -> dict:
    """Registra un lead interesado en el proyecto."""
    lead = {
        "telefono": telefono,
        "nombre": nombre,
        "interes": interes,
        "presupuesto": presupuesto,
        "fecha": datetime.now().isoformat(),
    }
    logger.info(f"Lead registrado: {lead}")
    return {"registrado": True}


LEADS_EXCEL_PATH = "data/leads.xlsx"

_HEADERS = ["Fecha", "Nombre", "Teléfono", "Email", "Apto de interés", "Habitaciones", "Temperatura", "Intención"]
_HEADER_COLOR = "1a3c5e"
_ROW_COLORS = ["FFFFFF", "EAF1F8"]


def exportar_leads_excel(leads: list[dict]) -> str:
    """Genera o sobreescribe data/leads.xlsx con todos los leads. Retorna la ruta."""
    wb = Workbook()
    ws = wb.active
    ws.title = _cfg().get("excel", {}).get("hoja_leads", "Leads")

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(fill_type="solid", fgColor=_HEADER_COLOR)
    header_align = Alignment(horizontal="center", vertical="center")

    for col_idx, header in enumerate(_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    ws.row_dimensions[1].height = 20

    for row_idx, lead in enumerate(leads, start=2):
        icono = ICONOS_TEMPERATURA.get(lead.get("temperatura", "").lower(), "")
        temp = f"{icono} {lead.get('temperatura', '').upper()}" if lead.get("temperatura") else ""
        fila = [
            lead.get("fecha", ""),
            lead.get("nombre", ""),
            lead.get("telefono", ""),
            lead.get("email", ""),
            lead.get("apto", ""),
            lead.get("habitaciones", ""),
            temp,
            lead.get("intencion", ""),
        ]
        fill_color = _ROW_COLORS[(row_idx - 2) % 2]
        row_fill = PatternFill(fill_type="solid", fgColor=fill_color)
        for col_idx, value in enumerate(fila, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = row_fill
            cell.alignment = Alignment(vertical="center")

    col_widths = [18, 25, 18, 30, 18, 14, 14, 14]
    for col_idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    os.makedirs("data", exist_ok=True)
    wb.save(LEADS_EXCEL_PATH)
    logger.info(f"Excel de leads actualizado: {len(leads)} leads")
    return LEADS_EXCEL_PATH


_VISITAS_EXCEL_PATH = "data/visitas.xlsx"
_VISITAS_HEADERS = ["Fecha visita", "Hora", "Nombre", "Teléfono", "Estado", "Notas", "Creado"]
_ESTADO_COLORS = {"completada": "C6EFCE", "cancelada": "FFC7CE", "confirmada": "DDEBF7"}

def exportar_visitas_excel(visitas: list[dict]) -> str:
    """Genera data/visitas.xlsx con todas las visitas. Retorna la ruta."""
    wb = Workbook()
    ws = wb.active
    ws.title = _cfg().get("excel", {}).get("hoja_visitas", "Visitas")

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(fill_type="solid", fgColor=_HEADER_COLOR)
    header_align = Alignment(horizontal="center", vertical="center")

    for col_idx, h in enumerate(_VISITAS_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
    ws.row_dimensions[1].height = 20

    for row_idx, v in enumerate(visitas, start=2):
        estado = v.get("estado", "confirmada")
        row_color = _ESTADO_COLORS.get(estado, "FFFFFF")
        row_fill = PatternFill(fill_type="solid", fgColor=row_color)
        fila = [
            v.get("fecha", ""),
            v.get("hora", ""),
            v.get("nombre", ""),
            v.get("telefono", ""),
            estado.upper(),
            v.get("notas", ""),
            (v.get("creado_at") or "")[:16],
        ]
        for col_idx, value in enumerate(fila, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = row_fill
            cell.alignment = Alignment(vertical="center")

    col_widths = [14, 8, 28, 18, 14, 35, 18]
    for col_idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    os.makedirs("data", exist_ok=True)
    wb.save(_VISITAS_EXCEL_PATH)
    logger.info(f"Excel de visitas generado: {len(visitas)} visitas")
    return _VISITAS_EXCEL_PATH


def registrar_visita(telefono: str, nombre: str, fecha: str, hora: str) -> dict:
    """Registra una solicitud de visita al proyecto."""
    visita = {
        "telefono": telefono,
        "nombre": nombre,
        "fecha_visita": fecha,
        "hora_visita": hora,
        "registrado": datetime.now().isoformat(),
    }
    logger.info(f"Visita agendada: {visita}")
    return {
        "confirmado": True,
        "mensaje": f"Visita agendada para el {fecha} a las {hora}."
    }


# ── Calculadora hipotecaria ───────────────────────────────────────────────────

def _fmt_cop(v: float) -> str:
    """Formatea un valor en pesos colombianos: $1.288.200.000"""
    return "$" + f"{v:,.0f}".replace(",", ".")


def calcular_hipoteca(precio_cop: float, entrada_pct: float, plazo_anos: int, tasa_anual_pct: float) -> str:
    """
    Calcula la cuota mensual estimada de un crédito hipotecario en Colombia.
    Retorna un resumen con todos los datos del cálculo.
    """
    entrada = precio_cop * (entrada_pct / 100)
    credito = precio_cop - entrada
    r = (tasa_anual_pct / 100) / 12
    n = int(plazo_anos) * 12
    if r == 0:
        cuota = credito / n
    else:
        cuota = credito * r * (1 + r) ** n / ((1 + r) ** n - 1)
    total = cuota * n
    intereses = total - credito
    return (
        f"Precio del apartamento: {_fmt_cop(precio_cop)}\n"
        f"Cuota inicial ({entrada_pct:.0f}%): {_fmt_cop(entrada)}\n"
        f"Monto del crédito: {_fmt_cop(credito)}\n"
        f"Tasa de interés: {tasa_anual_pct}% anual\n"
        f"Plazo: {plazo_anos} años ({n} cuotas)\n"
        f"CUOTA MENSUAL ESTIMADA: {_fmt_cop(cuota)}\n"
        f"Total pagado en {plazo_anos} años: {_fmt_cop(total)}\n"
        f"Total de intereses: {_fmt_cop(intereses)}"
    )
